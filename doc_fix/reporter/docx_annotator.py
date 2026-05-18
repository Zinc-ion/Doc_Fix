"""Write an annotated .docx review copy with highlights and comments."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.text.run import Run

from doc_fix.model import CheckIssue, CheckReport


COMMENT_AUTHOR = "Doc_Fix"
COMMENT_INITIALS = "DF"
HIGHLIGHT_COLOR = WD_COLOR_INDEX.YELLOW


def write_annotated_docx(report: CheckReport, output_path: Path) -> tuple[Path, tuple[str, ...]]:
    """Create a highlighted/commented copy of the target .docx."""

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document(report.input_docx_path)
    warnings: list[str] = []

    for issue in report.issues:
        runs = _issue_runs(document, issue)
        if not runs:
            warnings.append(f"{issue.code}: 未能在目标工作副本中定位可标注的文字或对象。")
            continue
        for run in runs:
            run.font.highlight_color = HIGHLIGHT_COLOR
        document.add_comment(
            runs,
            text=_comment_text(issue),
            author=COMMENT_AUTHOR,
            initials=COMMENT_INITIALS,
        )

    document.save(output_path)
    return output_path, tuple(warnings)


def _issue_runs(document, issue: CheckIssue) -> list[Run]:
    if issue.code.startswith("table.") and issue.table_index is not None:
        return _table_runs(document, issue.table_index) or _paragraph_runs(document, issue.paragraph_index)
    if issue.paragraph_index is not None:
        return _paragraph_range_runs(document, issue.paragraph_index, issue.end_paragraph_index)
    if issue.code.startswith("image.") and issue.paragraph_index is not None:
        return _paragraph_runs(document, issue.paragraph_index)
    return []


def _paragraph_range_runs(document, start_index: int, end_index: int | None) -> list[Run]:
    end_index = start_index if end_index is None else end_index
    if start_index < 0 or end_index < start_index:
        return []
    runs: list[Run] = []
    for paragraph_index in range(start_index, end_index + 1):
        runs.extend(_paragraph_runs(document, paragraph_index))
    return runs


def _paragraph_runs(document, paragraph_index: int | None) -> list[Run]:
    if paragraph_index is None or paragraph_index < 0 or paragraph_index >= len(document.paragraphs):
        return []
    return list(document.paragraphs[paragraph_index].runs)


def _table_runs(document, table_index: int) -> list[Run]:
    if table_index < 0 or table_index >= len(document.tables):
        return []
    table = document.tables[table_index]
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                runs = list(paragraph.runs)
                if runs and paragraph.text.strip():
                    return runs
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                if paragraph.runs:
                    return list(paragraph.runs)
    return []


def _comment_text(issue: CheckIssue) -> str:
    parts = [
        f"[{issue.severity.upper()}] {issue.code}",
        issue.message,
    ]
    if issue.locator:
        parts.append(f"定位：{issue.locator}")
    if issue.expected is not None or issue.actual is not None:
        parts.append(f"期望/实际：{_value(issue.expected)} / {_value(issue.actual)}")
    if issue.content_preview:
        parts.append(f"摘录：{issue.content_preview}")
    return "\n".join(parts)


def _value(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ")
