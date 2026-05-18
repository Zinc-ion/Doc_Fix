"""Extract document structure from .docx files."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from doc_fix.model import DocumentSnapshot, ImageSpec, ParagraphBlock, TableSpec


SECTION_HEADING_PATTERNS = (
    (re.compile(r"^第[一二三四五六七八九十百千万0-9]+部分\b.*"), 1),
    (re.compile(r"^[一二三四五六七八九十]+[、.．]\s*.+"), 2),
    (re.compile(r"^（[一二三四五六七八九十]+）\s*.+"), 3),
    (re.compile(r"^\([一二三四五六七八九十]+\)\s*.+"), 3),
    (re.compile(r"^[0-9]+[、.．]\s*.+"), 4),
)
TABLE_CAPTION_PATTERN = re.compile(r"^表\s*[0-9一二三四五六七八九十]+[\s、.．-]*(.+)?")
IMAGE_CAPTION_PATTERN = re.compile(r"^图\s*[0-9一二三四五六七八九十]+[\s、.．-]*(.+)?")


class DocxExtractor:
    """Build a lightweight document snapshot using python-docx and OOXML."""

    def extract(self, path: Path) -> DocumentSnapshot:
        document = Document(str(path))
        context = _StructureContext()
        paragraphs = []
        for index, paragraph in enumerate(document.paragraphs):
            text = paragraph.text.strip()
            style_name = paragraph.style.name if paragraph.style is not None else None
            heading_level = self._paragraph_heading_level(paragraph) or infer_chinese_heading_level(text)
            is_heading = heading_level is not None and bool(text)
            if is_heading:
                context.update_heading(heading_level, text)
            if TABLE_CAPTION_PATTERN.match(text):
                context.last_table_caption = text
            if IMAGE_CAPTION_PATTERN.match(text):
                context.last_image_caption = text
            if text:
                context.last_non_empty_text = text
            paragraphs.append(
                ParagraphBlock(
                    index=index,
                    text=text,
                    style_name=style_name,
                    heading_level=heading_level,
                    chapter_path=context.chapter_path,
                    is_heading=is_heading,
                    font_names=self._paragraph_font_names(paragraph),
                    font_sizes_pt=self._paragraph_font_sizes(paragraph),
                    alignment=self._alignment_name(paragraph.alignment),
                    first_line_indent_twips=self._length_twips(paragraph.paragraph_format.first_line_indent),
                    left_indent_twips=self._length_twips(paragraph.paragraph_format.left_indent),
                    right_indent_twips=self._length_twips(paragraph.paragraph_format.right_indent),
                    line_spacing=self._line_spacing(paragraph.paragraph_format.line_spacing),
                    space_before_twips=self._length_twips(paragraph.paragraph_format.space_before),
                    space_after_twips=self._length_twips(paragraph.paragraph_format.space_after),
                )
            )

        table_locations = self._locate_tables(document)
        tables = tuple(
            self._extract_table(index, table, table_locations.get(index))
            for index, table in enumerate(document.tables)
        )
        images = tuple(self._extract_images(document, paragraphs))

        page_width_twips = None
        page_content_width_twips = None
        if document.sections:
            section = document.sections[0]
            page_width_twips = int(section.page_width.twips)
            page_content_width_twips = int(
                section.page_width.twips
                - section.left_margin.twips
                - section.right_margin.twips
            )

        return DocumentSnapshot(
            source_path=str(path),
            paragraphs=paragraphs,
            tables=tables,
            images=images,
            page_width_twips=page_width_twips,
            page_content_width_twips=page_content_width_twips,
        )

    def _extract_table(self, index: int, table, location: _TableLocation | None = None) -> TableSpec:
        rows = len(table.rows)
        columns = len(table.columns)
        text = "\n".join(cell.text.strip() for row in table.rows for cell in row.cells if cell.text.strip())
        width_twips = self._table_width_twips(table)
        return TableSpec(
            index=index,
            rows=rows,
            columns=columns,
            text=text,
            chapter_path=location.chapter_path if location else None,
            chapter_table_index=location.chapter_table_index if location else None,
            caption=location.caption if location else None,
            nearby_heading=location.nearby_heading if location else None,
            paragraph_index=location.paragraph_index if location else None,
            width_twips=width_twips,
            has_borders=self._table_has_borders(table),
            alignment=self._table_alignment_name(table.alignment),
            style_name=table.style.name if table.style is not None else None,
            border_signature=self._table_border_signature(table),
            row_heights_twips=self._row_heights(table),
            column_widths_twips=self._column_widths(table),
            font_names=self._table_font_names(table),
            font_sizes_pt=self._table_font_sizes(table),
            cell_alignments=self._table_cell_alignments(table),
        )

    def _extract_images(self, document, paragraphs: list[ParagraphBlock]) -> list[ImageSpec]:
        images: list[ImageSpec] = []
        image_index = 0
        for paragraph_index, paragraph in enumerate(document.paragraphs):
            block = paragraphs[paragraph_index]
            alignment = self._alignment_name(paragraph.alignment)
            for drawing in paragraph._p.xpath(".//w:drawing"):
                inline = drawing.xpath(".//wp:inline")
                anchor = drawing.xpath(".//wp:anchor")
                ext_nodes = drawing.xpath(".//wp:extent")
                width = None
                height = None
                if ext_nodes:
                    width = self._safe_int(ext_nodes[0].get("cx"))
                    height = self._safe_int(ext_nodes[0].get("cy"))
                wrap = "anchor" if anchor else "inline" if inline else None
                images.append(
                    ImageSpec(
                        index=image_index,
                        width_emu=width,
                        height_emu=height,
                        paragraph_index=paragraph_index,
                        paragraph_alignment=alignment or wrap,
                        wrap_type=wrap,
                        chapter_path=block.chapter_path,
                        chapter_image_index=self._image_index_in_chapter(images, block.chapter_path) + 1,
                        caption=self._nearby_image_caption(paragraphs, paragraph_index),
                        nearby_heading=block.chapter_path,
                    )
                )
                image_index += 1
        return images

    def _locate_tables(self, document) -> dict[int, "_TableLocation"]:
        table_locations: dict[int, _TableLocation] = {}
        context = _StructureContext()
        chapter_counts: dict[str, int] = {}
        last_table_caption: str | None = None
        last_non_empty_text: str | None = None
        last_paragraph_index: int | None = None
        table_index = 0
        paragraph_index = 0

        for child in document.element.body.iterchildren():
            if child.tag.endswith("}p"):
                paragraph = Paragraph(child, document)
                text = paragraph.text.strip()
                heading_level = self._paragraph_heading_level(paragraph) or infer_chinese_heading_level(text)
                if heading_level is not None and text:
                    context.update_heading(heading_level, text)
                if TABLE_CAPTION_PATTERN.match(text):
                    last_table_caption = text
                if text:
                    last_non_empty_text = text
                    last_paragraph_index = paragraph_index
                paragraph_index += 1
                continue
            if child.tag.endswith("}tbl"):
                table = Table(child, document)
                chapter_path = context.chapter_path or "未能定位，需人工确认"
                count = chapter_counts.get(chapter_path, 0) + 1
                chapter_counts[chapter_path] = count
                table_locations[table_index] = _TableLocation(
                    chapter_path=chapter_path,
                    chapter_table_index=count,
                    caption=last_table_caption or first_non_empty_table_text(table),
                    nearby_heading=context.chapter_path or last_non_empty_text,
                    paragraph_index=last_paragraph_index,
                )
                last_table_caption = None
                table_index += 1
        return table_locations

    def _image_index_in_chapter(self, images: list[ImageSpec], chapter_path: str | None) -> int:
        return sum(1 for image in images if image.chapter_path == chapter_path)

    def _nearby_image_caption(self, paragraphs: list[ParagraphBlock], paragraph_index: int) -> str | None:
        for block in reversed(paragraphs[max(0, paragraph_index - 3) : paragraph_index + 1]):
            if IMAGE_CAPTION_PATTERN.match(block.text):
                return block.text
        return None

    def _table_width_twips(self, table) -> int | None:
        tbl_pr = table._tbl.tblPr
        if tbl_pr is None:
            return None
        tbl_w_nodes = tbl_pr.xpath("./w:tblW")
        if not tbl_w_nodes:
            return None
        tbl_w = tbl_w_nodes[0]
        width_type = tbl_w.get(qn("w:type"))
        if width_type not in (None, "dxa"):
            return None
        width = tbl_w.get(qn("w:w"))
        return self._safe_int(width)

    def _table_has_borders(self, table) -> bool | None:
        tbl_pr = table._tbl.tblPr
        if tbl_pr is None:
            return None
        tbl_borders = tbl_pr.xpath("./w:tblBorders")
        if not tbl_borders:
            return None
        return bool(list(tbl_borders[0]))

    def _table_border_signature(self, table) -> str | None:
        tbl_pr = table._tbl.tblPr
        if tbl_pr is None:
            return None
        tbl_borders = tbl_pr.xpath("./w:tblBorders")
        if not tbl_borders:
            return None
        parts: list[str] = []
        for border in tbl_borders[0]:
            tag = border.tag.rsplit("}", 1)[-1]
            val = border.get(qn("w:val"), "")
            size = border.get(qn("w:sz"), "")
            color = border.get(qn("w:color"), "")
            parts.append(f"{tag}:{val}:{size}:{color}")
        return "|".join(sorted(parts))

    def _row_heights(self, table) -> tuple[int | None, ...]:
        heights: list[int | None] = []
        for row in table.rows:
            heights.append(int(row.height.twips) if row.height is not None else None)
        return tuple(heights)

    def _column_widths(self, table) -> tuple[int | None, ...]:
        if not table.rows:
            return ()
        widths: list[int | None] = []
        for cell in table.rows[0].cells:
            widths.append(int(cell.width.twips) if cell.width is not None else None)
        return tuple(widths)

    def _table_font_names(self, table) -> tuple[str, ...]:
        names: set[str] = set()
        for cell in self._iter_cells(table):
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if run.font.name:
                        names.add(run.font.name)
        return tuple(sorted(names))

    def _table_font_sizes(self, table) -> tuple[float, ...]:
        sizes: set[float] = set()
        for cell in self._iter_cells(table):
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if run.font.size:
                        sizes.add(round(float(run.font.size.pt), 2))
        return tuple(sorted(sizes))

    def _paragraph_font_names(self, paragraph: Paragraph) -> tuple[str, ...]:
        names: set[str] = set()
        for run in paragraph.runs:
            if run.font.name:
                names.add(run.font.name)
            r_pr = run._r.rPr
            if r_pr is None or r_pr.rFonts is None:
                continue
            for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
                value = r_pr.rFonts.get(qn(f"w:{attr}"))
                if value:
                    names.add(value)
        return tuple(sorted(names))

    def _paragraph_font_sizes(self, paragraph: Paragraph) -> tuple[float, ...]:
        sizes: set[float] = set()
        for run in paragraph.runs:
            if run.font.size:
                sizes.add(round(float(run.font.size.pt), 2))
        return tuple(sorted(sizes))

    def _table_cell_alignments(self, table) -> tuple[str, ...]:
        alignments: set[str] = set()
        for cell in self._iter_cells(table):
            for paragraph in cell.paragraphs:
                alignment = self._alignment_name(paragraph.alignment)
                if alignment:
                    alignments.add(alignment)
        return tuple(sorted(alignments))

    def _iter_cells(self, table):
        for row in table.rows:
            for cell in row.cells:
                yield cell

    def _alignment_name(self, alignment) -> str | None:
        if alignment is None:
            return None
        names = {
            WD_ALIGN_PARAGRAPH.LEFT: "left",
            WD_ALIGN_PARAGRAPH.CENTER: "center",
            WD_ALIGN_PARAGRAPH.RIGHT: "right",
            WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
        }
        return names.get(alignment, str(alignment))

    def _length_twips(self, value) -> int | None:
        return int(value.twips) if value is not None else None

    def _line_spacing(self, value) -> float | None:
        if value is None:
            return None
        if hasattr(value, "twips"):
            return float(value.twips)
        try:
            return round(float(value), 4)
        except (TypeError, ValueError):
            return None

    def _table_alignment_name(self, alignment) -> str | None:
        if alignment is None:
            return None
        names = {
            WD_TABLE_ALIGNMENT.LEFT: "left",
            WD_TABLE_ALIGNMENT.CENTER: "center",
            WD_TABLE_ALIGNMENT.RIGHT: "right",
        }
        return names.get(alignment, str(alignment))

    def _paragraph_heading_level(self, paragraph) -> int | None:
        outline_nodes = paragraph._p.xpath("./w:pPr/w:outlineLvl")
        if outline_nodes:
            value = self._safe_int(outline_nodes[0].get(qn("w:val")))
            if value is not None:
                return value + 1

        style_name = paragraph.style.name if paragraph.style is not None else ""
        for prefix in ("Heading ", "标题 "):
            if style_name.startswith(prefix):
                return self._safe_int(style_name.removeprefix(prefix))
        return None

    def _safe_int(self, value: str | None) -> int | None:
        if value is None:
            return None
        try:
            return int(value)
        except ValueError:
            return None


class _StructureContext:
    def __init__(self) -> None:
        self._headings: dict[int, str] = {}
        self.last_table_caption: str | None = None
        self.last_image_caption: str | None = None
        self.last_non_empty_text: str | None = None

    @property
    def chapter_path(self) -> str | None:
        parts = [self._headings[level] for level in sorted(self._headings)]
        return " > ".join(parts) if parts else None

    def update_heading(self, level: int, text: str) -> None:
        self._headings[level] = text
        for existing_level in list(self._headings):
            if existing_level > level:
                del self._headings[existing_level]


class _TableLocation:
    def __init__(
        self,
        chapter_path: str,
        chapter_table_index: int,
        caption: str | None,
        nearby_heading: str | None,
        paragraph_index: int | None,
    ) -> None:
        self.chapter_path = chapter_path
        self.chapter_table_index = chapter_table_index
        self.caption = caption
        self.nearby_heading = nearby_heading
        self.paragraph_index = paragraph_index


def infer_chinese_heading_level(text: str) -> int | None:
    text = text.strip()
    if not text or TABLE_CAPTION_PATTERN.match(text) or IMAGE_CAPTION_PATTERN.match(text):
        return None
    if len(text) > 80:
        return None
    for pattern, level in SECTION_HEADING_PATTERNS:
        if pattern.match(text):
            return level
    return None


def first_non_empty_table_text(table) -> str | None:
    for row in table.rows:
        for cell in row.cells:
            text = " ".join(cell.text.split())
            if text:
                return text[:80]
    return None
