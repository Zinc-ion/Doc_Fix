from pathlib import Path
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from doc_fix.corrector import DocxCorrector
from doc_fix.extractor import DocxExtractor


def test_corrector_copies_paragraph_format_and_removes_bracket_remarks(tmp_path: Path) -> None:
    template_path = tmp_path / "template.docx"
    input_path = tmp_path / "input.docx"
    output_path = tmp_path / "input.corrected.docx"

    template = Document()
    template.add_paragraph("一、章节")
    reference = template.add_paragraph()
    reference.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    reference.paragraph_format.first_line_indent = Pt(24)
    reference.paragraph_format.left_indent = Pt(12)
    reference.paragraph_format.right_indent = Pt(6)
    reference.paragraph_format.line_spacing = 1.5
    reference.paragraph_format.space_before = Pt(6)
    reference.paragraph_format.space_after = Pt(12)
    reference_run = reference.add_run("模板正文")
    reference_run.font.name = "宋体"
    reference_run.font.size = Pt(14)
    reference_run.font.color.rgb = RGBColor(0x11, 0x22, 0x33)
    reference_run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    template.save(template_path)

    target = Document()
    target.add_paragraph("一、章节")
    paragraph = target.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run("目标正文【备注】")
    run.font.name = "黑体"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(0xAA, 0xBB, 0xCC)
    target.save(input_path)

    extractor = DocxExtractor()
    report = DocxCorrector().correct(
        template_path,
        input_path,
        output_path,
        extractor.extract(template_path),
        extractor.extract(input_path),
    )

    corrected = Document(output_path)
    corrected_paragraph = corrected.paragraphs[1]
    corrected_run = corrected_paragraph.runs[0]
    assert corrected_paragraph.text == "目标正文"
    assert corrected_paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert corrected_paragraph.paragraph_format.first_line_indent.pt == 24
    assert corrected_paragraph.paragraph_format.left_indent.pt == 12
    assert corrected_paragraph.paragraph_format.right_indent.pt == 6
    assert corrected_paragraph.paragraph_format.line_spacing == 1.5
    assert corrected_paragraph.paragraph_format.space_before.pt == 6
    assert corrected_paragraph.paragraph_format.space_after.pt == 12
    assert corrected_run.font.size.pt == 14
    assert corrected_run.font.color.rgb == RGBColor(0, 0, 0)
    assert corrected_run.font.highlight_color is None
    assert {action.code for action in report.actions} == {
        "format.color_normalized",
        "paragraph.format_aligned",
        "remarks.bracket_removed",
    }


def test_corrector_removes_multi_run_remarks_and_empty_body_paragraphs(tmp_path: Path) -> None:
    template_path = tmp_path / "template.docx"
    input_path = tmp_path / "input.docx"
    output_path = tmp_path / "input.corrected.docx"

    template = Document()
    template.add_paragraph().add_run("模板正文").font.size = Pt(12)
    template.save(template_path)

    target = Document()
    target.add_paragraph("【整段删除】")
    multi_run = target.add_paragraph()
    multi_run.add_run("保留【跨")
    multi_run.add_run("run】结束【第二个】")
    target.save(input_path)

    extractor = DocxExtractor()
    DocxCorrector().correct(
        template_path,
        input_path,
        output_path,
        extractor.extract(template_path),
        extractor.extract(input_path),
    )

    corrected = Document(output_path)
    assert [paragraph.text for paragraph in corrected.paragraphs] == ["保留结束"]


def test_corrector_keeps_out_of_scope_table_format_and_removes_table_remarks(tmp_path: Path) -> None:
    template_path = tmp_path / "template.docx"
    input_path = tmp_path / "input.docx"
    output_path = tmp_path / "input.corrected.docx"

    template = Document()
    template_table = template.add_table(rows=1, cols=1)
    template_paragraph = template_table.cell(0, 0).paragraphs[0]
    template_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    template_run = template_paragraph.add_run("模板单元格")
    template_run.font.name = "宋体"
    template_run.font.size = Pt(12)
    template_run.font.color.rgb = RGBColor(0, 0, 0)
    template.save(template_path)

    target = Document()
    target_table = target.add_table(rows=1, cols=1)
    target_cell = target_table.cell(0, 0)
    target_cell.text = "目标【表格备注】"
    target_cell.paragraphs[0].runs[0].font.size = Pt(16)
    target.save(input_path)

    extractor = DocxExtractor()
    report = DocxCorrector().correct(
        template_path,
        input_path,
        output_path,
        extractor.extract(template_path),
        extractor.extract(input_path),
    )

    corrected = Document(output_path)
    paragraph = corrected.tables[0].cell(0, 0).paragraphs[0]
    assert paragraph.text == "目标"
    assert paragraph.alignment is None
    assert paragraph.runs[0].font.size.pt == 16
    action_codes = {action.code for action in report.actions}
    assert "remarks.bracket_removed" in action_codes
    assert "format.color_normalized" in action_codes
    assert "table.text_format_aligned" not in action_codes


def test_corrector_normalizes_first_to_fifth_part_table_text_without_template(tmp_path: Path) -> None:
    template_path = tmp_path / "template.docx"
    input_path = tmp_path / "input.docx"
    output_path = tmp_path / "input.corrected.docx"

    template = Document()
    template_table = template.add_table(rows=1, cols=1)
    template_table.cell(0, 0).text = "模板表格"
    template_table.cell(0, 0).paragraphs[0].runs[0].font.size = Pt(18)
    template.save(template_path)

    target = Document()
    cover_table = target.add_table(rows=1, cols=1)
    cover_table.cell(0, 0).text = "封面表格 English"
    cover_table.cell(0, 0).paragraphs[0].runs[0].font.size = Pt(16)
    target.add_paragraph("第一部分  国内外现状及趋势分析")
    first_table = target.add_table(rows=1, cols=1)
    first_table.cell(0, 0).text = "第一部分表格 English"
    first_table.cell(0, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    first_table.cell(0, 0).paragraphs[0].runs[0].font.size = Pt(16)
    target.add_paragraph("第五部分  项目组织实施、保障措施及风险分析")
    fifth_table = target.add_table(rows=1, cols=1)
    fifth_table.cell(0, 0).text = "第五部分表格 English"
    fifth_table.cell(0, 0).paragraphs[0].runs[0].font.size = Pt(14)
    target.add_paragraph("第六部分  研究团队")
    sixth_table = target.add_table(rows=1, cols=1)
    sixth_table.cell(0, 0).text = "第六部分表格 English"
    sixth_table.cell(0, 0).paragraphs[0].runs[0].font.size = Pt(15)
    target.save(input_path)

    extractor = DocxExtractor()
    report = DocxCorrector().correct(
        template_path,
        input_path,
        output_path,
        extractor.extract(template_path),
        extractor.extract(input_path),
    )

    corrected = Document(output_path)
    cover_run = corrected.tables[0].cell(0, 0).paragraphs[0].runs[0]
    first_paragraph = corrected.tables[1].cell(0, 0).paragraphs[0]
    first_run = first_paragraph.runs[0]
    fifth_run = corrected.tables[2].cell(0, 0).paragraphs[0].runs[0]
    sixth_run = corrected.tables[3].cell(0, 0).paragraphs[0].runs[0]

    assert cover_run.font.size.pt == 16
    assert first_run.font.size.pt == 10.5
    assert fifth_run.font.size.pt == 10.5
    assert sixth_run.font.size.pt == 15
    assert first_run._r.rPr.rFonts.get(qn("w:eastAsia")) == "宋体"
    assert first_run._r.rPr.rFonts.get(qn("w:ascii")) == "Times New Roman"
    assert fifth_run._r.rPr.rFonts.get(qn("w:eastAsia")) == "宋体"
    assert fifth_run._r.rPr.rFonts.get(qn("w:ascii")) == "Times New Roman"
    assert first_paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert "table.text_format_aligned" in {action.code for action in report.actions}


def test_corrector_normalizes_package_xml_highlight_shading_and_colors(tmp_path: Path) -> None:
    template_path = tmp_path / "template.docx"
    input_path = tmp_path / "input.docx"
    output_path = tmp_path / "input.corrected.docx"

    template = Document()
    template.add_paragraph().add_run("模板正文").font.size = Pt(12)
    template.save(template_path)

    target = Document()
    paragraph = target.add_paragraph("带编号底纹的正文")
    paragraph.runs[0].font.color.rgb = RGBColor(0xFF, 0, 0)
    paragraph.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
    p_pr = paragraph._p.get_or_add_pPr()
    r_pr = OxmlElement("w:rPr")
    highlight = OxmlElement("w:highlight")
    highlight.set(qn("w:val"), "lightGray")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "00B050")
    r_pr.append(highlight)
    r_pr.append(color)
    r_pr.append(_shading("D9D9D9"))
    p_pr.append(r_pr)
    table = target.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格"
    table.cell(0, 0)._tc.get_or_add_tcPr().append(_shading("FFFF00"))
    target.save(input_path)

    extractor = DocxExtractor()
    report = DocxCorrector().correct(
        template_path,
        input_path,
        output_path,
        extractor.extract(template_path),
        extractor.extract(input_path),
    )

    with zipfile.ZipFile(output_path) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
    assert "<w:highlight" not in document_xml
    assert "<w:shd" not in document_xml
    assert 'w:val="FF0000"' not in document_xml
    assert 'w:val="00B050"' not in document_xml
    assert 'w:val="000000"' in document_xml
    assert "format.color_normalized" in {action.code for action in report.actions}


def test_corrector_body_fallback_uses_common_body_format_not_cover(tmp_path: Path) -> None:
    template_path = tmp_path / "template.docx"
    input_path = tmp_path / "input.docx"
    output_path = tmp_path / "input.corrected.docx"

    template = Document()
    cover = template.add_paragraph()
    cover_run = cover.add_run("国家重点研发计划")
    cover_run.font.name = "黑体"
    cover_run.font.size = Pt(22)
    template.add_paragraph("一、模板章节")
    first_body = template.add_paragraph()
    first_body.add_run("模板正文一").font.size = Pt(12)
    second_body = template.add_paragraph()
    second_body.add_run("模板正文二").font.size = Pt(12)
    template.save(template_path)

    target = Document()
    target.add_paragraph("二、目标章节")
    target_body = target.add_paragraph()
    target_body.add_run("目标正文").font.size = Pt(10.5)
    target.save(input_path)

    extractor = DocxExtractor()
    DocxCorrector().correct(
        template_path,
        input_path,
        output_path,
        extractor.extract(template_path),
        extractor.extract(input_path),
    )

    corrected = Document(output_path)
    assert corrected.paragraphs[1].runs[0].font.size.pt == 12


def test_corrector_no_chapter_paragraphs_match_by_text_before_fallback(tmp_path: Path) -> None:
    template_path = tmp_path / "template.docx"
    input_path = tmp_path / "input.docx"
    output_path = tmp_path / "input.corrected.docx"

    template = Document()
    cover = template.add_paragraph()
    cover.add_run("国家重点研发计划").font.size = Pt(22)
    date = template.add_paragraph()
    date.add_run("中华人民共和国科学技术部制").font.size = Pt(16)
    template.add_paragraph("一、模板章节")
    template.add_paragraph().add_run("模板正文").font.size = Pt(12)
    template.add_paragraph().add_run("更多正文").font.size = Pt(12)
    template.save(template_path)

    target = Document()
    target.add_paragraph().add_run("国家重点研发计划").font.size = Pt(9)
    target.add_paragraph().add_run("中华人民共和国科学技术部制").font.size = Pt(9)
    target.save(input_path)

    extractor = DocxExtractor()
    DocxCorrector().correct(
        template_path,
        input_path,
        output_path,
        extractor.extract(template_path),
        extractor.extract(input_path),
    )

    corrected = Document(output_path)
    assert corrected.paragraphs[0].runs[0].font.size.pt == 22
    assert corrected.paragraphs[1].runs[0].font.size.pt == 16


def _shading(fill: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    return shd
