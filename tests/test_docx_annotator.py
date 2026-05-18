from pathlib import Path
from zipfile import ZipFile

from docx import Document

from doc_fix.model import CheckIssue, CheckReport
from doc_fix.reporter import write_annotated_docx


def test_write_annotated_docx_adds_comments_and_highlights(tmp_path: Path) -> None:
    source_path = tmp_path / "input.converted.docx"
    output_path = tmp_path / "input.annotated.docx"
    document = Document()
    document.add_paragraph("第一段存在字数问题。")
    document.add_paragraph("第二段继续统计。")
    document.add_paragraph("图1 技术路线图")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格内容"
    document.save(source_path)
    report = CheckReport(
        template_path="template.doc",
        input_path="target.doc",
        template_docx_path="template.docx",
        input_docx_path=str(source_path),
        issues=(
            CheckIssue(
                code="word_count.exceeded",
                severity="error",
                message="字数超限",
                paragraph_index=0,
                end_paragraph_index=1,
                expected=3,
                actual=5,
            ),
            CheckIssue(
                code="table.columns_mismatch",
                severity="error",
                message="表格列数不一致",
                table_index=0,
            ),
            CheckIssue(
                code="image.width_mismatch",
                severity="warning",
                message="图片宽度不一致",
                paragraph_index=2,
            ),
        ),
    )

    annotated_path, warnings = write_annotated_docx(report, output_path)

    assert annotated_path == output_path
    assert warnings == ()
    with ZipFile(output_path) as package:
        names = set(package.namelist())
        document_xml = package.read("word/document.xml").decode("utf-8")
        comments_xml = package.read("word/comments.xml").decode("utf-8")
        rels_xml = package.read("word/_rels/document.xml.rels").decode("utf-8")
        content_types_xml = package.read("[Content_Types].xml").decode("utf-8")

    assert "word/comments.xml" in names
    assert "relationships/comments" in rels_xml
    assert "/word/comments.xml" in content_types_xml
    assert "commentRangeStart" in document_xml
    assert "commentRangeEnd" in document_xml
    assert "commentReference" in document_xml
    assert 'w:val="yellow"' in document_xml
    assert "word_count.exceeded" in comments_xml
    assert "table.columns_mismatch" in comments_xml
    assert "image.width_mismatch" in comments_xml


def test_write_annotated_docx_records_unlocated_issues(tmp_path: Path) -> None:
    source_path = tmp_path / "input.converted.docx"
    output_path = tmp_path / "input.annotated.docx"
    document = Document()
    document.add_paragraph("正文")
    document.save(source_path)
    report = CheckReport(
        template_path="template.doc",
        input_path="target.doc",
        template_docx_path="template.docx",
        input_docx_path=str(source_path),
        issues=(
            CheckIssue(
                code="word_count.anchor_missing",
                severity="warning",
                message="未能定位锚点",
            ),
        ),
    )

    annotated_path, warnings = write_annotated_docx(report, output_path)

    assert annotated_path.exists()
    assert len(warnings) == 1
    assert "word_count.anchor_missing" in warnings[0]


def test_write_annotated_docx_saves_clean_copy_without_issues(tmp_path: Path) -> None:
    source_path = tmp_path / "input.converted.docx"
    output_path = tmp_path / "input.annotated.docx"
    document = Document()
    document.add_paragraph("正文")
    document.save(source_path)
    report = CheckReport(
        template_path="template.doc",
        input_path="target.doc",
        template_docx_path="template.docx",
        input_docx_path=str(source_path),
    )

    annotated_path, warnings = write_annotated_docx(report, output_path)

    assert annotated_path.exists()
    assert warnings == ()
