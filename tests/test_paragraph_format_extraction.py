from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from doc_fix.extractor import DocxExtractor


def test_docx_extractor_reads_explicit_paragraph_format(tmp_path: Path) -> None:
    path = tmp_path / "formatted.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(24)
    paragraph.paragraph_format.left_indent = Pt(12)
    paragraph.paragraph_format.right_indent = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(12)
    run = paragraph.add_run("正文")
    run.font.name = "宋体"
    run.font.size = Pt(12)
    heading = document.add_heading("一、标题", level=1)
    heading.runs[0].font.name = "黑体"
    heading.runs[0].font.size = Pt(16)
    document.save(path)

    snapshot = DocxExtractor().extract(path)

    body = snapshot.paragraphs[0]
    assert body.font_names == ("宋体",)
    assert body.font_sizes_pt == (12.0,)
    assert body.alignment == "justify"
    assert body.first_line_indent_twips == 480
    assert body.left_indent_twips == 240
    assert body.right_indent_twips == 120
    assert body.line_spacing == 1.5
    assert body.space_before_twips == 120
    assert body.space_after_twips == 240

    extracted_heading = snapshot.paragraphs[1]
    assert extracted_heading.is_heading is True
    assert extracted_heading.heading_level == 1
    assert extracted_heading.font_names == ("黑体",)
    assert extracted_heading.font_sizes_pt == (16.0,)
