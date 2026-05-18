import json
import shutil
from pathlib import Path

from click.testing import CliRunner
from docx import Document
from docx.shared import Pt

from doc_fix.cli.main import main
from doc_fix.converter import ConversionResult, WordConverter


def test_correct_command_writes_outputs_and_keeps_original(tmp_path: Path, monkeypatch) -> None:
    template_path = tmp_path / "template.docx"
    input_path = tmp_path / "input.docx"
    out_dir = tmp_path / "out"

    template = Document()
    template.add_paragraph("项目简介 限10字以内")
    reference = template.add_paragraph("模板正文")
    reference.runs[0].font.size = Pt(12)
    template.save(template_path)

    target = Document()
    target.add_paragraph("项目简介 限10字以内")
    target.add_paragraph("目标正文【备注】")
    target.save(input_path)
    original_bytes = input_path.read_bytes()

    original_normalize = WordConverter.normalize_to_docx

    def fake_normalize(self, input_file: Path, output_dir: Path, label: str) -> ConversionResult:
        if input_file.suffix.lower() == ".doc":
            output_path = output_dir / f"{label}.converted.docx"
            shutil.copy2(output_dir / "input.corrected.docx", output_path)
            return ConversionResult(input_file, output_path, converted=True)
        return original_normalize(self, input_file, output_dir, label)

    def fake_export(self, docx_path: Path, output_path: Path) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(b"mock doc bytes")
        return output_path

    monkeypatch.setattr(WordConverter, "normalize_to_docx", fake_normalize)
    monkeypatch.setattr(WordConverter, "export_to_doc", fake_export)

    result = CliRunner().invoke(
        main,
        [
            "correct",
            "--template",
            str(template_path),
            "--input",
            str(input_path),
            "--out-dir",
            str(out_dir),
        ],
    )

    assert result.exit_code == 0, result.output
    assert input_path.read_bytes() == original_bytes
    assert (out_dir / "input.converted.docx").exists()
    assert (out_dir / "input.corrected.docx").exists()
    assert (out_dir / "input.corrected.doc").exists()
    assert (out_dir / "correction_report.json").exists()
    assert (out_dir / "corrected_check_report.json").exists()
    assert (out_dir / "corrected_doc_check_report.json").exists()

    corrected = Document(out_dir / "input.corrected.docx")
    assert corrected.paragraphs[1].text == "目标正文"

    correction_report = json.loads((out_dir / "correction_report.json").read_text(encoding="utf-8"))
    assert correction_report["corrected_doc_path"] == str(out_dir / "input.corrected.doc")
    assert correction_report["corrected_doc_export_error"] is None
    assert correction_report["corrected_check_report_path"] == str(out_dir / "corrected_check_report.json")
    assert correction_report["corrected_doc_check_report_path"] == str(out_dir / "corrected_doc_check_report.json")


def test_correct_command_records_doc_export_failure(tmp_path: Path, monkeypatch) -> None:
    template_path = tmp_path / "template.docx"
    input_path = tmp_path / "input.docx"
    out_dir = tmp_path / "out"

    template = Document()
    template.add_paragraph().add_run("模板正文").font.size = Pt(12)
    template.save(template_path)
    target = Document()
    target.add_paragraph("目标【备注】")
    target.save(input_path)

    def fake_export(self, docx_path: Path, output_path: Path) -> Path:
        from doc_fix.converter import ConversionError

        raise ConversionError("Word unavailable")

    monkeypatch.setattr(WordConverter, "export_to_doc", fake_export)

    result = CliRunner().invoke(
        main,
        [
            "correct",
            "--template",
            str(template_path),
            "--input",
            str(input_path),
            "--out-dir",
            str(out_dir),
        ],
    )

    assert result.exit_code == 0, result.output
    correction_report = json.loads((out_dir / "correction_report.json").read_text(encoding="utf-8"))
    assert correction_report["corrected_doc_path"] is None
    assert "Word unavailable" in correction_report["corrected_doc_export_error"]
    assert correction_report["corrected_doc_check_report_path"] is None
